"""Inviting a whole team at once, without inventing a second way to invite anyone.

Two operations, and the split is the point. `classify` writes nothing and tells the
administrator what would happen to each address; `invite_many` does it. An admin who
pastes eighty addresses gets to see the six that are already members and the two that are
misspelt *before* seventy-two people receive mail, which is the difference between a bulk
tool and a bulk mistake.

**Every actual invitation still goes through `invite_employee`.** That function owns the
rank ceiling, the already-a-member refusal, the live-invitation conflict, the org-less
token index, the audit row and the mail. A second insert path here would be a second
place for those to be wrong, and the one that skipped `outranks` would be a privilege
escalation with a CSV attached.

**Each row gets a savepoint, and that is not defensive habit.** Postgres aborts a
transaction at its first error, and one of the two refusals here IS an error: a live
invitation for the address is an `IntegrityError` from the INSERT, after which every
statement on that connection fails until the transaction ends. So one stale address in a
re-pasted list would discard every invitation after it while the batch reported success —
the trap this codebase already documents for `record_failure`. A nested transaction per
row makes a refusal a result rather than the end of the batch.

(The already-a-member refusal is not that: it comes from a clean SELECT and raises in
Python, leaving the connection healthy. Both arrive as `Conflict`, which is precisely why
the distinction has to be written down rather than inferred from the exception type.)
"""

from __future__ import annotations

import asyncio
import csv
import io
import logging
import re
from dataclasses import dataclass
from enum import StrEnum
from typing import Final
from uuid import UUID

from jutsu_core.errors import Conflict, PermissionDenied, ValidationFailed
from jutsu_core.rbac import Role, outranks
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from jutsu_api.config import Settings
from jutsu_api.email import EmailMessage, EmailSender
from jutsu_api.invitations import invite_employee
from jutsu_api.security import Principal

logger = logging.getLogger("jutsu.invitations")

__all__ = [
    "MAX_BULK_ROWS",
    "MAX_XLSX_BYTES",
    "ROLE_TITLE_MAX",
    "BulkOutcome",
    "BulkRow",
    "ClassifiedRow",
    "InviteResult",
    "classify",
    "invite_many",
    "parse_csv",
    "parse_pasted",
    "parse_xlsx",
]

#: How many addresses one request may carry.
#:
#: Two hundred is comfortably more than a team lands at once, and a larger import is
#: several batches — which is also how an administrator wants to review one.
MAX_BULK_ROWS: Final = 200

#: How long a role TITLE may be, everywhere it is handled.
#:
#: The parsers truncate to it, the request model bounds by it, and `invite_employee`
#: truncates to it again on the way to the database. Three places agreeing is not
#: redundancy: when the parser did NOT truncate, a spreadsheet with a 200-character title
#: produced a preview row the browser posted straight back into a `max_length=128` field,
#: and Pydantic rejected the WHOLE batch at the request boundary over one cell.
ROLE_TITLE_MAX: Final = 128

#: How many messages are in flight at once during the delivery pass.
#:
#: `SmtpEmailSender` opens a fresh connection per message — connect, STARTTLS, LOGIN,
#: DATA — which is most of a second against a real provider. Two hundred of those in
#: sequence is nearly three minutes inside one request, so they overlap. Five rather than
#: fifty because submission endpoints rate-limit concurrent logins, and a bulk import that
#: trips the provider's throttle has failed in the way that is hardest to explain.
_DELIVERY_CONCURRENCY: Final = 5

#: Deliberately permissive, because the authority on an address is the mailbox that
#: answers it. This rejects what is obviously not an address — no @, whitespace inside,
#: a missing dot in the domain — and leaves the rest to delivery. A stricter pattern here
#: would reject valid addresses and teach administrators to distrust the preview.
_EMAIL: Final = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

#: What separates one CONTACT from the next: a comma, a semicolon or a line break.
#:
#: Whitespace is deliberately not here. `Ada Lovelace <ada@example.com>` is one contact
#: containing two spaces, so splitting on whitespace first would destroy it; the spaces
#: inside a segment are dealt with below, after the bracketed form has been recognised.
_SEGMENTS: Final = re.compile(r"[,;\n\r]+")

#: A quoted display name, removed before segmenting.
#:
#: `"Babbage, C" <cb@example.com>` is one contact whose *name* contains the character that
#: otherwise separates contacts, and RFC 5322 quotes it for exactly that reason. Stripping
#: the quoted span first is what stops the comma inside it splitting one person into two.
#: Only double quotes: an apostrophe is a legal character in a local part (`o'brien@…`),
#: so treating `'` as a delimiter would corrupt real addresses.
_QUOTED_NAME: Final = re.compile(r'"[^"\n]*"')

#: The address inside `<…>`, which is what every mail client copies around the name.
#:
#: **Linear by construction, and that is the point.** The previous version tried to match
#: the display name too, with a greedy `[^"'<>,;\n]*` that could match empty and overlapped
#: the `\s*` beside it — so a 64 000-character paste with no `<` in it made the engine retry
#: every split point at every start position and hung the event loop. It also swallowed any
#: bare address sitting before a bracketed contact on the same line, silently dropping that
#: person from the import. Matching only the bracketed part removes both: the character
#: classes here are disjoint, so there is nothing to backtrack over.
_ANGLED_ADDRESS: Final = re.compile(r"<\s*([^<>\s,;]+)\s*>")


class BulkOutcome(StrEnum):
    """What `classify` says will happen, and what `invite_many` reports happened.

    The two share a vocabulary on purpose: the preview an administrator approved and the
    result they get back read the same way, so a row that changed between them is
    visible rather than merely different.
    """

    #: Will be invited, or was.
    READY = "ready"
    SENT = "sent"
    #: Already an active member of this organisation. Not an error — the usual result of
    #: re-pasting last month's list.
    ALREADY_MEMBER = "already_member"
    #: A live invitation is already outstanding for this address.
    ALREADY_INVITED = "already_invited"
    #: The same address appears earlier in this input. Only the first is acted on.
    DUPLICATE = "duplicate"
    #: Not a usable address.
    INVALID_EMAIL = "invalid_email"
    #: Not a role in the catalogue.
    INVALID_ROLE = "invalid_role"
    #: The inviter does not outrank the requested role. Named plainly here because the
    #: administrator is entitled to know why their own import refused a row — the vague
    #: wording in `invite_employee` protects against probing by a *caller*, and this
    #: caller already holds `member:invite` over their own organisation.
    ROLE_TOO_HIGH = "role_too_high"
    #: The invitation could not be created or sent. The row can be retried.
    FAILED = "failed"


@dataclass(frozen=True, slots=True)
class BulkRow:
    """One requested invitation, as the administrator supplied it."""

    email: str
    role: Role = Role.MEMBER
    role_title: str | None = None
    #: The literal text of a `role` cell that named nothing in the catalogue.
    #:
    #: Carried rather than resolved, because the alternative is silently seating somebody
    #: as a Member because their spreadsheet said "Manager". `classify` turns this into an
    #: `invalid_role` row that names the word it did not understand, so the administrator
    #: fixes the file instead of discovering the wrong access later.
    role_error: str | None = None


@dataclass(frozen=True, slots=True)
class ClassifiedRow:
    """A row and what would happen to it, with the reason a person can act on."""

    email: str
    role: Role
    role_title: str | None
    outcome: BulkOutcome
    #: One sentence, written for the administrator reading the preview. Never carries a
    #: token, and never says anything about an address outside this organisation.
    detail: str


@dataclass(frozen=True, slots=True)
class InviteResult:
    rows: list[ClassifiedRow]

    @property
    def sent(self) -> int:
        return sum(1 for row in self.rows if row.outcome is BulkOutcome.SENT)

    @property
    def failed(self) -> int:
        return sum(1 for row in self.rows if row.outcome is BulkOutcome.FAILED)


def parse_pasted(raw: str, *, default_role: Role = Role.MEMBER) -> list[BulkRow]:
    """Addresses from whatever a person pasted.

    Newlines, commas, semicolons and stray whitespace all separate; `Ada Lovelace
    <ada@example.com>` yields the address, because that is what a mail client copies.
    Everything unrecognised is kept rather than dropped, so `classify` can show it as
    invalid instead of the administrator wondering where a line went.
    """
    rows: list[BulkRow] = []
    # Quoted names go first, so a comma inside one cannot split a contact in half.
    for segment in _SEGMENTS.split(_QUOTED_NAME.sub(" ", raw or "")):
        segment = segment.strip()
        if not segment:
            continue

        if _ANGLED_ADDRESS.search(segment):
            # **Walked in order, because a segment can hold both forms.**
            # `ada@example.com Grace Hopper <grace@example.com>` is two people, and taking
            # only the bracketed one drops Ada as surely as the greedy pattern this
            # replaced did. Between the brackets, only tokens containing `@` survive —
            # everything else is a display name, and keeping it would produce a row per
            # word.
            position = 0
            for match in _ANGLED_ADDRESS.finditer(segment):
                rows.extend(_loose_addresses(segment[position : match.start()], default_role))
                rows.append(BulkRow(email=match.group(1), role=default_role))
                position = match.end()
            rows.extend(_loose_addresses(segment[position:], default_role))
            continue

        # No brackets anywhere in this segment, so whitespace is the only thing left that
        # can separate two addresses — and anything here that is not one survives to be
        # shown as invalid rather than being silently dropped.
        for token in segment.split():
            cleaned = token.strip("<>\"'")
            if cleaned:
                rows.append(BulkRow(email=cleaned, role=default_role))
    return rows


def _loose_addresses(fragment: str, role: Role) -> list[BulkRow]:
    """Bare addresses sitting beside a `Name <addr>` contact.

    Only tokens carrying an `@`. In this position the alternative is a display name, and
    a preview that turns "Grace Hopper" into two invalid rows is one nobody reads.
    """
    found: list[BulkRow] = []
    for token in fragment.split():
        cleaned = token.strip("<>\"',;")
        if "@" in cleaned:
            found.append(BulkRow(email=cleaned, role=role))
    return found


def parse_csv(raw: str, *, default_role: Role = Role.MEMBER) -> list[BulkRow]:
    """Rows from a CSV whose header names the columns.

    `email` is required. `role` and `role_title` are honoured when present; every other
    column is ignored rather than rejected, because an export from an HR system carries
    twenty columns this product has no field for, and refusing the file over them would
    make the feature unusable with real data.

    A file with no recognisable header is treated as a single column of addresses, which
    is what a list saved out of a spreadsheet actually looks like.
    """
    text_io = io.StringIO(raw.lstrip("﻿"))
    try:
        sample = raw[:2048]
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t") if sample.strip() else csv.excel
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(text_io, dialect)
    try:
        records = [row for row in reader if any(cell.strip() for cell in row)]
    except csv.Error as exc:
        # **The iteration raises, not just the sniff.** `csv.field_size_limit()` defaults
        # to 131 072 characters, comfortably inside the megabyte this route accepts, and a
        # single over-long field makes `csv.reader` raise mid-iteration — outside the
        # `except` above, which only ever covered the dialect sniff. Unhandled, that left
        # the route answering 500 with a traceback for what is simply a bad file.
        raise ValidationFailed("That file could not be read as a spreadsheet.") from exc

    if not records:
        return []

    header = [cell.strip().lower().replace(" ", "_") for cell in records[0]]
    if "email" not in header:
        # No header: every non-empty first cell is an address. `default_role` is passed
        # here for the same reason the header branch and `parse_pasted` pass it —
        # omitting it silently seated a whole headerless import as Member whatever the
        # administrator had chosen in the dropdown.
        return [
            BulkRow(email=row[0].strip(), role=default_role)
            for row in records
            if row and row[0].strip()
        ]

    index = {name: position for position, name in enumerate(header)}
    rows: list[BulkRow] = []
    for record in records[1:]:
        email = _cell(record, index, "email")
        if not email:
            # **Reported, not dropped.** A record with content but no address — a short
            # row, or a file whose addresses are in a column this does not read — used to
            # vanish, so the preview's total silently disagreed with the file the person
            # was looking at. Carrying the first cell that does have content makes the row
            # `invalid_email` with something recognisable in it.
            stray = next((cell.strip() for cell in record if cell.strip()), "")
            if stray:
                rows.append(BulkRow(email=stray, role=default_role))
            continue
        written_role = _cell(record, index, "role")
        try:
            role = Role(written_role.lower().replace(" ", "_")) if written_role else default_role
        except ValueError:
            rows.append(BulkRow(email=email, role=default_role, role_error=written_role))
            continue
        rows.append(
            BulkRow(
                email=email,
                role=role,
                role_title=_cell(record, index, "role_title")[:ROLE_TITLE_MAX] or None,
            )
        )
    return rows


def _cell(record: list[str], index: dict[str, int], name: str) -> str:
    """One named column of one row, or "" when the file does not have it.

    A free function rather than a closure over the loop: the closure was correct — it was
    called inside the iteration that bound it — but a reader has to prove that, and a
    later edit that stored it would silently read the last row for every row.
    """
    position = index.get(name)
    if position is None or position >= len(record):
        return ""
    return record[position].strip()


#: The largest workbook this will open, in bytes of the file as uploaded.
#:
#: A spreadsheet of two hundred addresses is a few kilobytes. Anything approaching a
#: megabyte is either the wrong file or an attempt to make the parser work, and refusing
#: it by size is the check that runs before any untrusted bytes reach a parser at all.
MAX_XLSX_BYTES: Final = 1_000_000

#: The largest the archive may claim to expand to, across all its members.
#:
#: **`MAX_XLSX_BYTES` bounds the COMPRESSED file, which is not a bound on the work.** A
#: `.xlsx` is a zip, and XML compresses about a thousand to one — so a one-megabyte upload
#: can declare a gigabyte of contents. `read_only=True` does not save us either: before it
#: defers anything, `load_workbook` reads the content types, the theme, the stylesheet and
#: the ENTIRE shared-string table into a Python list, so `_MAX_SHEET_ROWS` never gets the
#: chance to apply. This is checked against the zip's own directory before openpyxl is
#: handed the bytes.
_MAX_XLSX_UNCOMPRESSED: Final = 40_000_000

#: How many members the archive may hold. A real workbook has a few dozen.
_MAX_XLSX_MEMBERS: Final = 256

#: How many rows are read out of the sheet before it stops.
#:
#: A `.xlsx` is a zip, so a small upload can decompress into a very large sheet — and
#: `read_only=True` streams rather than materialising, which means the bound has to be the
#: loop's, not the file's. Generous enough to reach `MAX_BULK_ROWS` plus a header and
#: whatever blank rows a spreadsheet carries at the bottom.
_MAX_SHEET_ROWS: Final = 2_000


def _refuse_a_zip_bomb(data: bytes) -> None:
    """Read the archive's own directory before any of it is expanded.

    The central directory records each member's uncompressed length, so the cost of the
    file can be known without paying it. An attacker can of course lie in that header —
    but then the declared size no longer matches the stream, and `zipfile` raises during
    the read, which `parse_xlsx` already turns into a refusal. What this closes is the
    honest bomb: a small archive that truthfully declares a gigabyte, which openpyxl would
    otherwise load into a Python list of shared strings before streaming anything.
    """
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > _MAX_XLSX_MEMBERS:
                raise ValidationFailed("That file could not be read as a spreadsheet.")
            if sum(member.file_size for member in members) > _MAX_XLSX_UNCOMPRESSED:
                raise ValidationFailed(
                    "That file is too large. Export the addresses and try again."
                )
    except ValidationFailed:
        raise
    except Exception as exc:
        raise ValidationFailed("That file could not be read as a spreadsheet.") from exc


def parse_xlsx(data: bytes, *, default_role: Role = Role.MEMBER) -> list[BulkRow]:
    """Rows from the workbook an HR team already has.

    Read under four bounds, because this is the one input here that is a container format
    rather than text: the file is refused above `MAX_XLSX_BYTES` before it is opened; only
    the first worksheet is read; only `_MAX_SHEET_ROWS` rows are pulled from it; and
    `read_only=True` streams the sheet instead of building the whole thing in memory.
    `data_only=True` takes the stored result of a formula rather than the formula — an
    address is a value, and evaluating a spreadsheet is not something this should do.

    The cells are then handed to `parse_csv`, so a workbook and a CSV agree about what a
    header means and there is one place that decides it.
    """
    if len(data) > MAX_XLSX_BYTES:
        raise ValidationFailed("That file is too large. Export the addresses and try again.")

    _refuse_a_zip_bomb(data)

    # Imported here rather than at module scope: openpyxl pulls in its own XML machinery,
    # and every API process would pay for it at startup for a route most of them never
    # serve.
    from openpyxl import load_workbook

    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            return []
        table: list[list[str]] = []
        # **Inside the guard, not after it.** `read_only=True` parses the sheet lazily, so
        # a workbook whose header opens cleanly and whose rows are corrupt raises HERE —
        # and while this loop sat outside the `except`, that answered 500 with a traceback
        # for a file the administrator only needed to be told was unreadable.
        for index, cells in enumerate(sheet.iter_rows(values_only=True)):
            if index >= _MAX_SHEET_ROWS:
                break
            table.append(["" if cell is None else str(cell).strip() for cell in cells])
    except ValidationFailed:
        raise
    except Exception as exc:
        # Anything at all: a truncated zip, a `.xls`, a renamed PDF, corrupt sheet XML.
        # One refusal, because the distinctions are not ones the administrator can act on
        # differently.
        raise ValidationFailed("That file could not be read as a spreadsheet.") from exc
    finally:
        if workbook is not None:
            workbook.close()

    # Back through the CSV reader rather than a second header implementation. Quoting the
    # cells is what keeps a value containing a comma — a role title, typically — from
    # becoming two columns on the way through.
    buffer = io.StringIO()
    csv.writer(buffer, lineterminator="\n").writerows(table)
    return parse_csv(buffer.getvalue(), default_role=default_role)


def _normalise(email: str) -> str:
    return email.strip().lower()


async def classify(
    session: AsyncSession, *, actor: Principal, rows: list[BulkRow]
) -> list[ClassifiedRow]:
    """What would happen to each row. Writes nothing.

    Two lookups for the whole batch rather than two per row: an import of two hundred
    addresses would otherwise be four hundred round trips to answer a question the
    administrator has not yet agreed to act on.
    """
    if len(rows) > MAX_BULK_ROWS:
        raise ValidationFailed(
            f"That is {len(rows)} addresses. Import up to {MAX_BULK_ROWS} at a time."
        )

    candidates = [_normalise(row.email) for row in rows]
    wanted = [address for address in candidates if _EMAIL.fullmatch(address)]

    members: set[str] = set()
    invited: set[str] = set()
    if wanted:
        members = {
            str(value).lower()
            for value in (
                await session.execute(
                    text(
                        "SELECT lower(email) FROM users "
                        "WHERE lower(email) = ANY(:addresses) AND status = 'active'"
                    ),
                    {"addresses": wanted},
                )
            ).scalars()
        }
        invited = {
            str(value).lower()
            for value in (
                await session.execute(
                    text(
                        "SELECT lower(email) FROM invitations "
                        "WHERE lower(email) = ANY(:addresses) "
                        "AND accepted_at IS NULL AND revoked_at IS NULL AND expires_at > now()"
                    ),
                    {"addresses": wanted},
                )
            ).scalars()
        }

    seen: set[str] = set()
    classified: list[ClassifiedRow] = []
    for row in rows:
        address = _normalise(row.email)

        if not _EMAIL.fullmatch(address):
            outcome, detail = BulkOutcome.INVALID_EMAIL, "Not an email address."
        elif address in seen:
            outcome, detail = BulkOutcome.DUPLICATE, "Appears earlier in this list."
        elif row.role_error:
            outcome, detail = (
                BulkOutcome.INVALID_ROLE,
                f"'{row.role_error}' is not a role. Choose one from the list.",
            )
        elif address in members:
            outcome, detail = BulkOutcome.ALREADY_MEMBER, "Already in this organisation."
        elif address in invited:
            outcome, detail = BulkOutcome.ALREADY_INVITED, "Has an invitation waiting."
        elif not outranks(actor.role, row.role):
            outcome, detail = (
                BulkOutcome.ROLE_TOO_HIGH,
                "You cannot invite someone at that level of access.",
            )
        else:
            outcome, detail = BulkOutcome.READY, "Will be invited."

        seen.add(address)
        classified.append(
            ClassifiedRow(
                email=address,
                role=row.role,
                role_title=row.role_title,
                outcome=outcome,
                detail=detail,
            )
        )
    return classified


async def invite_many(
    session: AsyncSession,
    *,
    actor: Principal,
    rows: list[BulkRow],
    settings: Settings,
    sender: EmailSender,
) -> InviteResult:
    """Invite every row that can be invited, and report what happened to the rest.

    **A refused row must not take the batch with it.** `invite_employee` raises
    `Conflict` twice over, and the two are not alike: an address that is already a member
    fails a SELECT and raises in Python, while an address that already has a live
    invitation fails the INSERT against `uq_invitations_org_email_live` — a genuine
    Postgres error, after which the transaction is aborted and every later statement
    fails too. Without a savepoint per row that second case silently discards every
    invitation after it and still reports success. `begin_nested` makes each row its own
    unit: the failure rolls back to the savepoint and the loop continues.

    Re-running the same list is therefore safe and is the expected way to retry: the rows
    that succeeded come back as `already_invited`, and the ones that failed are tried
    again.

    **Delivery happens after the rows exist, not inside them.** `invite_employee` awaits
    the transport, and `SmtpEmailSender` opens a connection per message — so inviting a
    department one row at a time would spend minutes in a single request. Each row is
    given a `_Collected` transport instead; the messages it captures are delivered
    together at the end, `_DELIVERY_CONCURRENCY` at a time.

    That trade has one consequence and it is handled rather than accepted: a message that
    cannot be delivered no longer rolls its row back, because the savepoint is long
    released. The invitation is **revoked** instead, which reaches the same end state — no
    live invitation, no live token, and the address free to be invited again by the retry.
    """
    classified = await classify(session, actor=actor, rows=rows)
    results: list[ClassifiedRow] = []
    #: position in `results` → (invitation id, the messages that row produced)
    pending: dict[int, tuple[UUID, list[EmailMessage]]] = {}

    for row, verdict in zip(rows, classified, strict=True):
        if verdict.outcome is not BulkOutcome.READY:
            results.append(verdict)
            continue

        collected = _Collected()
        try:
            async with session.begin_nested():
                issued = await invite_employee(
                    session,
                    actor=actor,
                    email=verdict.email,
                    role=row.role,
                    settings=settings,
                    sender=collected,
                    role_title=row.role_title,
                )
        except Conflict as conflict:
            # Raced with another administrator, or with an earlier row of this very
            # batch. Both are "someone got there first", which is not a failure the
            # person needs to act on.
            results.append(
                _with(verdict, BulkOutcome.ALREADY_INVITED, str(conflict) or "Already invited.")
            )
        except PermissionDenied as denied:
            results.append(_with(verdict, BulkOutcome.ROLE_TOO_HIGH, str(denied)))
        except Exception as failure:
            # **The class name, and nothing else — not the message and not `exc_info`.**
            # SQLAlchemy renders bound parameters into its exception text, and the
            # parameter here is the address being invited: a traceback would put the
            # customer's mailing list in the log aggregator one import at a time (§4.9).
            # `hide_parameters=True` covers the production engine, but a log line must
            # not depend on an engine flag set somewhere else to stay clean.
            logger.warning(
                "%s", {"event": "bulk_invitation_row_failed", "error": type(failure).__name__}
            )
            results.append(
                _with(verdict, BulkOutcome.FAILED, "Could not be invited. Try this row again.")
            )
        else:
            pending[len(results)] = (issued.invitation_id, collected.messages)
            results.append(_with(verdict, BulkOutcome.SENT, "Invitation sent."))

    await _deliver(session, sender=sender, pending=pending, results=results)
    return InviteResult(rows=results)


class _Collected:
    """An `EmailSender` that keeps the message instead of delivering it.

    It exists so `invite_employee` stays the only thing that creates an invitation while
    the delivery it performs is lifted out of the per-row savepoint and batched. The
    messages carry live tokens in `EmailMessage.secrets`, which is `repr=False` and never
    logged — this object holds them for the length of one request and nothing else.
    """

    def __init__(self) -> None:
        self.messages: list[EmailMessage] = []

    async def send(self, message: EmailMessage) -> None:
        self.messages.append(message)


async def _deliver(
    session: AsyncSession,
    *,
    sender: EmailSender,
    pending: dict[int, tuple[UUID, list[EmailMessage]]],
    results: list[ClassifiedRow],
) -> None:
    """Send the collected messages, and revoke the invitations whose message did not go.

    Bounded concurrency rather than a `gather` over everything: a submission endpoint
    throttles concurrent logins, and two hundred at once is how an import gets the whole
    batch refused by the provider instead of one row refused by us.
    """
    if not pending:
        return

    limit = asyncio.Semaphore(_DELIVERY_CONCURRENCY)

    async def deliver_one(messages: list[EmailMessage]) -> bool:
        async with limit:
            try:
                for message in messages:
                    await sender.send(message)
            except Exception:
                # No address and no subject: §4.9, and the customer list must not end up
                # in the log aggregator one bulk import at a time.
                logger.warning("%s", {"event": "bulk_invitation_delivery_failed"})
                return False
        return True

    positions = list(pending)
    delivered = await asyncio.gather(*(deliver_one(pending[position][1]) for position in positions))

    undeliverable = [
        pending[position][0] for position, ok in zip(positions, delivered, strict=True) if not ok
    ]
    if undeliverable:
        # Revoked, not deleted. The row is the audit trail of an attempt that was made,
        # and the partial unique index only covers live invitations — so revoking is
        # exactly what lets the administrator press retry on that address.
        await session.execute(
            text("UPDATE invitations SET revoked_at = now() WHERE id = ANY(:ids)"),
            {"ids": undeliverable},
        )

    for position, ok in zip(positions, delivered, strict=True):
        if not ok:
            results[position] = _with(
                results[position],
                BulkOutcome.FAILED,
                "The invitation could not be emailed. Try this row again.",
            )


def _with(row: ClassifiedRow, outcome: BulkOutcome, detail: str) -> ClassifiedRow:
    return ClassifiedRow(
        email=row.email,
        role=row.role,
        role_title=row.role_title,
        outcome=outcome,
        detail=detail,
    )
